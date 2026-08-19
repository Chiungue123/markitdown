#!/usr/bin/env python3
"""Generate sample study materials so you can preview the whole pipeline.

Creates a small Biology 101 lecture as a PDF handout and a PowerPoint deck, so you
can run the converter against realistic input without supplying your own files.

USAGE:
    python study-tools/demo/make_sample_materials.py [OUTPUT_DIR]

    OUTPUT_DIR defaults to ~/Desktop/StudyMaterials/Demo — so the next step is:
        studyprep Demo

DEPENDENCIES (demo only):
    pip install reportlab python-pptx
    # python-pptx is already included if you installed markitdown[all] / [pptx];
    # only reportlab (for the PDF) is extra.
"""
import os
import sys

# Default output: a "Demo" subject inside the study hub, so `studyprep Demo` works.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
HUB = os.environ.get("STUDYPREP_HUB", os.path.join(os.path.expanduser("~"), "Desktop", "StudyMaterials"))
DEFAULT_OUT = os.path.join(HUB, "Demo")


def make_pdf(out_dir: str) -> str:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, ListFlowable, ListItem,
        )
    except ImportError:
        print("Skipping PDF: reportlab not installed (pip install reportlab).")
        return ""

    styles = getSampleStyleSheet()
    h1, h2, body = styles["Heading1"], styles["Heading2"], styles["BodyText"]
    path = os.path.join(out_dir, "photosynthesis-handout.pdf")
    doc = SimpleDocTemplate(path, pagesize=letter, title="Photosynthesis — Handout")
    doc.build([
        Paragraph("Biology 101 — Photosynthesis (Lecture Handout)", h1),
        Paragraph("Overview", h2),
        Paragraph("Photosynthesis is the process by which plants, algae, and some "
                  "bacteria convert light energy into chemical energy stored in "
                  "glucose. It is the foundation of nearly all food chains on Earth.", body),
        Paragraph("The Overall Equation", h2),
        Paragraph("6 CO2 + 6 H2O + light energy -> C6H12O6 + 6 O2", body),
        Paragraph("Two Stages", h2),
        ListFlowable([
            ListItem(Paragraph("<b>Light-dependent reactions</b> (thylakoid membranes): "
                               "capture light, split water, produce ATP and NADPH.", body)),
            ListItem(Paragraph("<b>Calvin cycle</b> (stroma): use ATP and NADPH to fix "
                               "CO2 into glucose.", body)),
        ], bulletType="bullet"),
        Paragraph("Key Terms", h2),
        ListFlowable([
            ListItem(Paragraph("<b>Chlorophyll</b>: green pigment that absorbs light.", body)),
            ListItem(Paragraph("<b>Stomata</b>: leaf pores for CO2 in and O2 out.", body)),
            ListItem(Paragraph("<b>ATP</b>: energy currency from the light reactions.", body)),
        ], bulletType="bullet"),
    ])
    return path


def make_pptx(out_dir: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        print("Skipping PPTX: python-pptx not installed (pip install python-pptx).")
        return ""

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Introduction to Photosynthesis"
    title.placeholders[1].text = "Biology 101 — Week 4 Lecture Slides"

    def bullets(heading, points):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        tf = slide.placeholders[1].text_frame
        tf.text = points[0]
        for p in points[1:]:
            tf.add_paragraph().text = p

    bullets("What Is Photosynthesis?", [
        "Converts light energy into chemical energy (glucose)",
        "Occurs in chloroplasts of plant cells",
        "Basis of almost every food chain on Earth",
    ])
    bullets("The Equation", [
        "6 CO2 + 6 H2O + light -> C6H12O6 + 6 O2",
        "Reactants: carbon dioxide, water, light",
        "Products: glucose (stored energy) and oxygen",
    ])
    bullets("Two Stages", [
        "Light-dependent reactions: produce ATP and NADPH",
        "Calvin cycle: fixes CO2 into glucose using ATP/NADPH",
        "Stage 1 in thylakoids; stage 2 in the stroma",
    ])
    bullets("Exam Tips", [
        "Memorize the balanced equation",
        "Know where each stage occurs",
        "Explain the role of chlorophyll",
    ])
    path = os.path.join(out_dir, "photosynthesis-slides.pptx")
    prs.save(path)
    return path


def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUT
    os.makedirs(out_dir, exist_ok=True)
    created = [p for p in (make_pdf(out_dir), make_pptx(out_dir)) if p]
    if created:
        print("Generated sample materials in", out_dir + ":")
        for p in created:
            print("  -", os.path.basename(p))
        print("\nNext: run  studyprep " + os.path.basename(out_dir))
    else:
        print("Nothing generated. Install demo dependencies:")
        print("  pip install reportlab python-pptx")
        sys.exit(1)


if __name__ == "__main__":
    main()
